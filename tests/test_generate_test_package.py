from __future__ import annotations

import importlib.util
import json
import tempfile
import unittest
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "测试规则" / "scripts" / "generate_test_package.py"
SPEC = importlib.util.spec_from_file_location("generate_test_package", SCRIPT)
assert SPEC and SPEC.loader
GENERATOR = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(GENERATOR)


class GenerateTestPackageTests(unittest.TestCase):
    def load_example(self) -> dict:
        path = ROOT / "测试规则" / "templates" / "test-manifest.example.json"
        return json.loads(path.read_text(encoding="utf-8"))

    def test_generate_and_check_package(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "input.json"
            output = root / "archive"
            source.write_text(json.dumps(self.load_example(), ensure_ascii=False), encoding="utf-8")

            GENERATOR.generate(source, output)
            GENERATOR.check(output)

            archived = json.loads((output / "source" / "manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(4, len(archived["artifacts"]))
            self.assertTrue((output / "证据").is_dir())
            for relative in archived["artifacts"]:
                self.assertTrue((output / relative).is_file())

    def test_wap_is_normalized_to_h5(self) -> None:
        data = self.load_example()
        data["cases"][0]["platform"] = "WAP"
        cases, _ = GENERATOR.validate(data)
        self.assertEqual("H5", cases[0]["platform"])

    def test_covered_item_requires_case_key(self) -> None:
        data = self.load_example()
        data["coverage"][0]["case_keys"] = []
        with self.assertRaises(SystemExit):
            GENERATOR.validate(data)

    def test_evidence_path_must_use_display_id(self) -> None:
        data = self.load_example()
        data["cases"][0]["execution"] = {
            "status": "passed",
            "note": "ok",
            "evidence_ids": ["ev-1"],
        }
        data["evidence"] = [{
            "evidence_id": "ev-1",
            "path": "证据/draft-core/save.json",
            "case_keys": ["module-core-flow"],
            "proves": "保存成功",
        }]
        with self.assertRaises(SystemExit):
            GENERATOR.validate(data)

    def test_status_colors_and_evidence_list(self) -> None:
        data = self.load_example()
        data["cases"].append({
            "case_key": "module-defect",
            "module": "所属模块",
            "platform": "PC",
            "keyword": "缺陷",
            "priority": 2,
            "name": "缺陷条",
            "precondition": "已登录。",
            "steps": ["操作失败现场。"],
            "expects": ["应成功。"],
            "timeout_seconds": 30,
            "execution": {
                "status": "suspected_defect",
                "note": "报错",
                "evidence_ids": ["ev-fail"],
            },
        })
        data["cases"][0]["execution"] = {
            "status": "passed",
            "note": "通过",
            "evidence_ids": ["ev-pass"],
        }
        data["coverage"].append({
            "coverage_id": "A-002",
            "group": "A",
            "checkpoint": "缺陷检查点",
            "source": "需求",
            "status": "covered",
            "case_keys": ["module-defect"],
            "reason": "",
        })
        data["evidence"] = [
            {
                "evidence_id": "ev-pass",
                "path": "证据/TC_所属模块_001/TC_所属模块_001_通过.json",
                "case_keys": ["module-core-flow"],
                "proves": "核心流程通过",
            },
            {
                "evidence_id": "ev-fail",
                "path": "证据/TC_所属模块_002/TC_所属模块_002_失败-白框.json",
                "case_keys": ["module-defect"],
                "proves": "白框",
            },
        ]
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "input.json"
            output = root / "archive"
            (output / "证据" / "TC_所属模块_001").mkdir(parents=True)
            (output / "证据" / "TC_所属模块_002").mkdir(parents=True)
            (output / "证据" / "TC_所属模块_001" / "TC_所属模块_001_通过.json").write_text("{}", encoding="utf-8")
            (output / "证据" / "TC_所属模块_002" / "TC_所属模块_002_失败-白框.json").write_text("{}", encoding="utf-8")
            source.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
            GENERATOR.generate(source, output)

            from openpyxl import load_workbook
            book = load_workbook(output / "所属产品-功能测试用例-执行状态清单.xlsx")
            sheet = book.worksheets[0]
            self.assertTrue(str(sheet.cell(2, 5).fill.fgColor.rgb or "").upper().endswith("C6EFCE"))
            self.assertTrue(str(sheet.cell(3, 5).fill.fgColor.rgb or "").upper().endswith("FFC7CE"))
            self.assertIn("证据/TC_所属模块_001/TC_所属模块_001_通过.json", str(sheet.cell(2, 7).value))
            book.close()
            listing = (output / "证据" / "实测证据清单.md").read_text(encoding="utf-8")
            self.assertIn("TC_所属模块_001", listing)
            self.assertIn("TC_所属模块_002", listing)
            self.assertIn("已实测通过", listing)
            self.assertIn("疑似缺陷", listing)


if __name__ == "__main__":
    unittest.main()
