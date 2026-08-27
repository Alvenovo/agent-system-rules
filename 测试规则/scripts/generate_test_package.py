#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import tempfile
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font


HEADERS = [
    "所属模块", "用例编号", "关键词", "优先级", "用例名称", "前置条件",
    "步骤", "预期", "测试结果（留空）", "用例备注", "用例类型", "适用阶段",
]
STATUSES = {
    "passed": "已实测通过",
    "partial": "部分验证",
    "blocked_environment": "当前环境受阻",
    "blocked_dependency": "前置依赖受阻",
    "suspected_defect": "疑似缺陷",
    "not_run": "未执行",
    "not_applicable": "不适用",
}
PLATFORMS = {"PC", "Android", "iOS", "H5", "MiniProgram"}
COVERAGE_STATUSES = {"covered", "blocked", "not_applicable"}


def fail(message: str) -> None:
    raise SystemExit(f"校验失败：{message}")


def read_json(path: Path) -> dict[str, Any]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        fail(f"无法读取 {path}: {exc}")
    if not isinstance(data, dict):
        fail("manifest 顶层必须是对象")
    return data


def required_text(value: Any, field: str) -> str:
    text = str(value or "").strip()
    if not text:
        fail(f"{field} 不能为空")
    return text


def safe_name(value: str) -> str:
    value = re.sub(r'[\\/:*?"<>|]', "", value).strip().replace(" ", "_")
    return value or "未命名"


def normalize_requirement_id(value: str) -> str:
    value = value.strip()
    match = re.fullmatch(r"([A-Za-z]+)_?(\d+)", value)
    return f"{match.group(1)}_{match.group(2)}" if match else value


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def atomic_write_text(path: Path, content: str, encoding: str = "utf-8") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    os.close(fd)
    temp = Path(temp_name)
    try:
        temp.write_text(content, encoding=encoding, newline="")
        temp.replace(path)
    finally:
        temp.unlink(missing_ok=True)


def atomic_save_workbook(workbook: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    temp = Path(temp_name)
    try:
        workbook.save(temp)
        temp.replace(path)
    finally:
        temp.unlink(missing_ok=True)


def numbered(items: list[Any]) -> str:
    return "\r\n".join(f"{index}. {str(item).strip()}" for index, item in enumerate(items, 1))


def validate(data: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, str]]:
    if data.get("schema_version") != 1:
        fail("schema_version 必须为 1")
    required_text(data.get("project"), "project")
    requirement = data.get("requirement")
    if not isinstance(requirement, dict):
        fail("requirement 必须是对象")
    for field in ("id", "name", "owner", "date"):
        required_text(requirement.get(field), f"requirement.{field}")

    materials = data.get("materials")
    if not isinstance(materials, list) or not materials:
        fail("materials 必须是非空数组")
    for index, material in enumerate(materials):
        if not isinstance(material, dict):
            fail(f"materials[{index}] 必须是对象")
        required_text(material.get("source"), f"materials[{index}].source")
        if material.get("status") not in {"read", "blocked"}:
            fail(f"materials[{index}].status 必须为 read 或 blocked")
        if material["status"] == "blocked":
            required_text(material.get("reason"), f"materials[{index}].reason")

    cases = data.get("cases")
    if not isinstance(cases, list) or not cases:
        fail("cases 必须是非空数组")
    keys: set[str] = set()
    counters: defaultdict[str, int] = defaultdict(int)
    display_ids: dict[str, str] = {}
    normalized: list[dict[str, Any]] = []
    for index, case in enumerate(cases):
        if not isinstance(case, dict):
            fail(f"cases[{index}] 必须是对象")
        key = required_text(case.get("case_key"), f"cases[{index}].case_key")
        if key in keys:
            fail(f"case_key 重复：{key}")
        keys.add(key)
        module = required_text(case.get("module"), f"cases[{index}].module")
        for field in ("keyword", "name", "precondition"):
            required_text(case.get(field), f"cases[{index}].{field}")
        steps = case.get("steps")
        expects = case.get("expects")
        if not isinstance(steps, list) or not steps or not isinstance(expects, list):
            fail(f"{key} 的 steps/expects 必须是非空数组")
        if len(steps) != len(expects):
            fail(f"{key} 的步骤与预期数量不一致")
        if any(not str(item).strip() for item in [*steps, *expects]):
            fail(f"{key} 的步骤或预期存在空项")
        priority = case.get("priority")
        if priority not in {1, 2, 3, 4}:
            fail(f"{key}.priority 必须为 1/2/3/4")
        platform = "H5" if case.get("platform") == "WAP" else case.get("platform")
        if platform not in PLATFORMS:
            fail(f"{key}.platform 必须为 PC/Android/iOS/H5/MiniProgram；WAP 会归一为 H5")
        execution = case.get("execution", {"status": "not_run", "note": "", "evidence_ids": []})
        if not isinstance(execution, dict) or execution.get("status") not in STATUSES:
            fail(f"{key}.execution.status 无效")
        if not isinstance(execution.get("evidence_ids", []), list):
            fail(f"{key}.execution.evidence_ids 必须是数组")
        timeout = case.get("timeout_seconds", 30)
        if not isinstance(timeout, int) or timeout <= 0:
            fail(f"{key}.timeout_seconds 必须是正整数")
        counters[module] += 1
        display_ids[key] = f"TC_{module}_{counters[module]:03d}"
        normalized.append({**case, "platform": platform, "execution": execution})

    coverage = data.get("coverage")
    if not isinstance(coverage, list) or not coverage:
        fail("coverage 必须是非空数组")
    coverage_ids: set[str] = set()
    for index, item in enumerate(coverage):
        if not isinstance(item, dict):
            fail(f"coverage[{index}] 必须是对象")
        coverage_id = required_text(item.get("coverage_id"), f"coverage[{index}].coverage_id")
        if coverage_id in coverage_ids:
            fail(f"coverage_id 重复：{coverage_id}")
        coverage_ids.add(coverage_id)
        if item.get("group") not in {"A", "B", "C"}:
            fail(f"{coverage_id}.group 必须为 A/B/C")
        required_text(item.get("checkpoint"), f"{coverage_id}.checkpoint")
        status = item.get("status")
        if status not in COVERAGE_STATUSES:
            fail(f"{coverage_id}.status 无效")
        case_keys = item.get("case_keys", [])
        if not isinstance(case_keys, list) or any(key not in keys for key in case_keys):
            fail(f"{coverage_id}.case_keys 包含不存在的用例")
        if status == "covered" and not case_keys:
            fail(f"{coverage_id} 已覆盖但没有 case_key")
        if status != "covered":
            required_text(item.get("reason"), f"{coverage_id}.reason")

    evidence = data.get("evidence", [])
    if not isinstance(evidence, list):
        fail("evidence 必须是数组")
    evidence_ids: set[str] = set()
    for index, item in enumerate(evidence):
        if not isinstance(item, dict):
            fail(f"evidence[{index}] 必须是对象")
        evidence_id = required_text(item.get("evidence_id"), f"evidence[{index}].evidence_id")
        if evidence_id in evidence_ids:
            fail(f"evidence_id 重复：{evidence_id}")
        evidence_ids.add(evidence_id)
        relative = Path(required_text(item.get("path"), f"{evidence_id}.path"))
        if relative.is_absolute() or ".." in relative.parts:
            fail(f"{evidence_id}.path 必须是归档目录内相对路径")
        item_keys = item.get("case_keys")
        if not isinstance(item_keys, list) or not item_keys or any(key not in keys for key in item_keys):
            fail(f"{evidence_id}.case_keys 无效")
        required_text(item.get("proves"), f"{evidence_id}.proves")

    for case in normalized:
        for evidence_id in case["execution"].get("evidence_ids", []):
            if evidence_id not in evidence_ids:
                fail(f"{case['case_key']} 引用了不存在的证据 {evidence_id}")
    return normalized, display_ids


def style_sheet(sheet: Any) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(
            60, max(12, max(len(str(cell.value or "")) for cell in column) + 2)
        )


def artifact_names(data: dict[str, Any]) -> dict[str, str]:
    project = safe_name(required_text(data.get("project"), "project"))
    requirement = data["requirement"]
    requirement_id = safe_name(normalize_requirement_id(required_text(requirement.get("id"), "requirement.id")))
    requirement_name = safe_name(required_text(requirement.get("name"), "requirement.name"))
    return {
        "csv": f"{project}-功能测试用例-byAI.csv",
        "cases_xlsx": f"{project}-_{requirement_id}_{requirement_name}.xlsx",
        "status_xlsx": f"{project}-功能测试用例-执行状态清单.xlsx",
        "process_md": f"{project}-功能测试用例-执行过程.md",
    }


def generate(source: Path, output: Path) -> None:
    data = read_json(source)
    cases, display_ids = validate(data)
    output.mkdir(parents=True, exist_ok=True)
    (output / "source").mkdir(exist_ok=True)
    (output / "证据").mkdir(exist_ok=True)
    names = artifact_names(data)

    previous_manifest = output / "source" / "manifest.json"
    if previous_manifest.exists():
        previous = read_json(previous_manifest)
        for relative in previous.get("artifacts", {}):
            if relative not in names.values():
                old = output / relative
                if old.is_file():
                    old.unlink()

    csv_path = output / names["csv"]
    with tempfile.NamedTemporaryFile(
        "w", encoding="utf-8-sig", newline="", delete=False, dir=output, prefix=".cases.", suffix=".csv"
    ) as stream:
        writer = csv.writer(stream, lineterminator="\r\n")
        writer.writerow(HEADERS)
        for case in cases:
            writer.writerow([
                case["module"], display_ids[case["case_key"]],
                f"{case['module']}_{case['keyword']}", case["priority"], case["name"],
                case["precondition"], numbered(case["steps"]), numbered(case["expects"]),
                "", "", "功能测试", "功能测试阶段",
            ])
        temp_csv = Path(stream.name)
    temp_csv.replace(csv_path)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "测试用例"
    sheet.append(HEADERS)
    for case in cases:
        sheet.append([
            case["module"], display_ids[case["case_key"]],
            f"{case['module']}_{case['keyword']}", case["priority"], case["name"],
            case["precondition"], numbered(case["steps"]), numbered(case["expects"]),
            "", "", "功能测试", "功能测试阶段",
        ])
    style_sheet(sheet)
    coverage_sheet = workbook.create_sheet("覆盖表")
    coverage_sheet.append(["覆盖编号", "分组", "检查点", "来源", "状态", "用例", "原因"])
    for item in data["coverage"]:
        coverage_sheet.append([
            item["coverage_id"], item["group"], item["checkpoint"], item.get("source", ""),
            item["status"], ", ".join(item.get("case_keys", [])), item.get("reason", ""),
        ])
    style_sheet(coverage_sheet)
    atomic_save_workbook(workbook, output / names["cases_xlsx"])

    status_book = Workbook()
    status_sheet = status_book.active
    status_sheet.title = "执行状态"
    status_sheet.append(["case_key", "用例编号", "平台", "用例名称", "执行状态", "实测说明", "证据"])
    for case in cases:
        execution = case["execution"]
        status_sheet.append([
            case["case_key"], display_ids[case["case_key"]], case["platform"], case["name"],
            STATUSES[execution["status"]], execution.get("note", ""),
            ", ".join(execution.get("evidence_ids", [])),
        ])
    style_sheet(status_sheet)
    atomic_save_workbook(status_book, output / names["status_xlsx"])

    lines = [
        f"# {data['project']}测试执行过程", "",
        f"- 需求：{data['requirement']['id']} {data['requirement']['name']}",
        f"- 负责人：{data['requirement']['owner']}",
        f"- 日期：{data['requirement']['date']}", "",
        "## 材料核对", "",
    ]
    for material in data["materials"]:
        suffix = f"；{material.get('reason', '')}" if material["status"] == "blocked" else ""
        lines.append(f"- [{material['status']}] {material['source']}{suffix}")
    lines.extend(["", "## 覆盖表", "", "| 编号 | 组 | 检查点 | 状态 | 用例/原因 |", "|---|---|---|---|---|"])
    for item in data["coverage"]:
        mapping = ", ".join(item.get("case_keys", [])) or item.get("reason", "")
        lines.append(f"| {item['coverage_id']} | {item['group']} | {item['checkpoint']} | {item['status']} | {mapping} |")
    lines.extend(["", "## 用例", ""])
    for case in cases:
        lines.extend([
            f"### {display_ids[case['case_key']]} {case['name']}",
            f"- case_key：`{case['case_key']}`",
            f"- 平台：{case['platform']}",
            f"- 前置条件：{case['precondition']}",
            f"- 状态：{STATUSES[case['execution']['status']]}",
            f"- 说明：{case['execution'].get('note', '') or '无'}",
            "",
        ])
        for index, (step, expect) in enumerate(zip(case["steps"], case["expects"]), 1):
            lines.append(f"{index}. 操作：{step}")
            lines.append(f"   预期：{expect}")
        lines.append("")
    atomic_write_text(output / names["process_md"], "\n".join(lines) + "\n")

    artifacts = {relative: sha256(output / relative) for relative in names.values()}
    archived = {**data, "artifacts": artifacts}
    atomic_write_text(
        previous_manifest,
        json.dumps(archived, ensure_ascii=False, indent=2) + "\n",
    )
    check(output)
    print(f"ok: generated {len(cases)} cases in {output}")


def check(output: Path) -> None:
    manifest_path = output / "source" / "manifest.json"
    data = read_json(manifest_path)
    cases, _ = validate(data)
    artifacts = data.get("artifacts")
    if not isinstance(artifacts, dict) or len(artifacts) != 4:
        fail("artifacts 必须记录四件套")
    for relative, expected_hash in artifacts.items():
        path = output / relative
        if not path.is_file():
            fail(f"缺少产物：{relative}")
        if sha256(path) != expected_hash:
            fail(f"产物哈希不一致：{relative}")
    names = artifact_names(data)
    with (output / names["csv"]).open("r", encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    if not rows or rows[0] != HEADERS or len(rows) - 1 != len(cases):
        fail("CSV 表头或条数不一致")
    for key in ("cases_xlsx", "status_xlsx"):
        book = load_workbook(output / names[key], read_only=True, data_only=True)
        sheet = book.worksheets[0]
        if sheet.max_row - 1 != len(cases):
            fail(f"{names[key]} 条数不一致")
        book.close()
    print(f"ok: package valid ({len(cases)} cases)")


def main() -> None:
    parser = argparse.ArgumentParser(description="从统一 manifest 生成或校验测试四件套")
    parser.add_argument("source", nargs="?", type=Path, help="源 manifest.json")
    parser.add_argument("output", nargs="?", type=Path, help="需求归档目录")
    parser.add_argument("--check", dest="check_dir", type=Path, help="校验已有需求归档目录")
    args = parser.parse_args()
    if args.check_dir:
        if args.source or args.output:
            fail("--check 不能与生成参数同时使用")
        check(args.check_dir.resolve())
        return
    if not args.source or not args.output:
        parser.error("生成时需要 <manifest.json> <归档目录>")
    generate(args.source.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
